import os
import tempfile
import unittest

import pandas as pd
from docx import Document
from PIL import Image

from generador_docx import crear_informe_docx, guardar_informe
from generador_pdf import generar_pdf_desde_docx
from lector_datos import calcular_puntuaciones_directas, leer_datos_excel
from reglas_psicometricas import obtener_puntuaciones


REPO_DIR = os.path.dirname(os.path.abspath(__file__))
SAMPLE_XLSX = os.path.join(REPO_DIR, 'excel ejemplo.xlsx')


class PipelineTests(unittest.TestCase):
    @staticmethod
    def _count_pdf_pages(path):
        with open(path, 'rb') as handle:
            return handle.read().count(b'/Type /Page')

    def test_sample_workbook_end_to_end(self):
        datos = leer_datos_excel(SAMPLE_XLSX)
        resultados = calcular_puntuaciones_directas(datos)
        clasificaciones = obtener_puntuaciones(resultados)

        self.assertIn('ACS', resultados['available_tests'])
        self.assertIn('CPT', resultados['available_tests'])
        self.assertIn('PT_ANT_A', resultados)
        self.assertIn('CPT_VAR_condicion', clasificaciones)
        self.assertIn('DUALTASK_PSV_cuando_concurrencia', clasificaciones)

        with tempfile.TemporaryDirectory() as tmpdir:
            ruta_docx = os.path.join(tmpdir, 'informe.docx')
            ruta_pdf = os.path.join(tmpdir, 'informe.pdf')
            doc = crear_informe_docx(resultados, clasificaciones, resultados['nombre_completo'], REPO_DIR)
            guardar_informe(doc, ruta_docx)
            self.assertTrue(os.path.exists(ruta_docx))
            self.assertTrue(generar_pdf_desde_docx(ruta_docx, ruta_pdf, verbose=False))
            self.assertTrue(os.path.exists(ruta_pdf))

    def test_detecta_pruebas_sustitutivas_y_ausentes(self):
        info = pd.DataFrame([{
            'datetime': '2026-01-01 10:00',
            'sub_num': 'Caso Sustituto',
            'age': 10,
        }])
        d2 = pd.DataFrame([
            {'row': 1, 'letter_num': 1, 'selected': True, 'timestamp': 100, 'target': 'si'},
            {'row': 1, 'letter_num': 2, 'selected': False, 'timestamp': 0, 'target': 'no'},
        ])
        five_digits = pd.DataFrame([
            {
                'part': 2, 'trial_type': 'experimental', 'contour': '1', 'content': '1',
                'discrepancy': 'no', 'response_given': '1', 'correct_response': '1',
                'correct': 'yes', 'TR': 1000,
            },
            {
                'part': 3, 'trial_type': 'experimental', 'contour': '2', 'content': '2',
                'discrepancy': 'no', 'response_given': '2', 'correct_response': '2',
                'correct': 'yes', 'TR': 1100,
            },
            {
                'part': 4, 'trial_type': 'experimental', 'contour': '3', 'content': '3',
                'discrepancy': 'yes', 'response_given': '3', 'correct_response': '3',
                'correct': 'yes', 'TR': 1200,
            },
        ])

        with tempfile.TemporaryDirectory() as tmpdir:
            ruta_xlsx = os.path.join(tmpdir, 'sustituto.xlsx')
            with pd.ExcelWriter(ruta_xlsx) as writer:
                info.to_excel(writer, sheet_name='info', index=False)
                d2.to_excel(writer, sheet_name='D2', index=False)
                five_digits.to_excel(writer, sheet_name='FiveDigits', index=False)

            datos = leer_datos_excel(ruta_xlsx)
            self.assertEqual(datos['display_names']['CPT'], 'D2')
            self.assertEqual(datos['display_names']['FourFigures'], 'FiveDigits')
            self.assertFalse(datos['test_presence']['DUALTASK'].present)

            resultados = calcular_puntuaciones_directas(datos)
            clasificaciones = obtener_puntuaciones(resultados)
            doc = crear_informe_docx(resultados, clasificaciones, 'Caso Sustituto', REPO_DIR)
            texto = "\n".join(p.text for p in doc.paragraphs)
            self.assertIn('D2 - prueba de rendimiento continuo', texto)
            self.assertIn('FiveDigits - control y flexibilidad cognitiva', texto)
            self.assertNotIn('Dual Task - prueba multitarea de atención dividida', texto)

    def test_tabla_y_pt_provisional_se_generan(self):
        datos = leer_datos_excel(SAMPLE_XLSX)
        resultados = calcular_puntuaciones_directas(datos)
        clasificaciones = obtener_puntuaciones(resultados)
        doc = crear_informe_docx(resultados, clasificaciones, resultados['nombre_completo'], REPO_DIR)
        self.assertGreaterEqual(len(doc.tables), 1)
        table_text = "\n".join(cell.text for row in doc.tables[0].rows for cell in row.cells)
        self.assertIn('Disponible', table_text)
        self.assertIn('PT', table_text)
        self.assertIsInstance(resultados['PT_DUALTASK_A'], int)

    def test_pdf_export_con_imagen_y_ruta_sin_directorio(self):
        datos = leer_datos_excel(SAMPLE_XLSX)
        resultados = calcular_puntuaciones_directas(datos)
        clasificaciones = obtener_puntuaciones(resultados)

        with tempfile.TemporaryDirectory() as tmpdir:
            ruta_docx = os.path.join(tmpdir, 'con_imagen.docx')
            ruta_docx_sin = os.path.join(tmpdir, 'sin_imagen.docx')
            guardar_informe(
                crear_informe_docx(resultados, clasificaciones, resultados['nombre_completo'], tmpdir),
                ruta_docx_sin,
            )
            self.assertTrue(generar_pdf_desde_docx(ruta_docx_sin, os.path.join(tmpdir, 'sin_imagen.pdf'), verbose=False))

            Image.new('RGB', (20, 20), 'white').save(os.path.join(tmpdir, 'grafico_CPT_final.png'))
            guardar_informe(
                crear_informe_docx(resultados, clasificaciones, resultados['nombre_completo'], tmpdir),
                ruta_docx,
            )

            cwd = os.getcwd()
            try:
                os.chdir(tmpdir)
                self.assertTrue(generar_pdf_desde_docx(ruta_docx, 'solo_nombre.pdf', verbose=False))
                self.assertTrue(os.path.exists(os.path.join(tmpdir, 'solo_nombre.pdf')))
                self.assertGreater(
                    self._count_pdf_pages(os.path.join(tmpdir, 'solo_nombre.pdf')),
                    self._count_pdf_pages(os.path.join(tmpdir, 'sin_imagen.pdf')),
                )
            finally:
                os.chdir(cwd)


if __name__ == '__main__':
    unittest.main()
